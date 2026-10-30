"""Generate a FAKE attendance workbook with the same layout as the real SAP export.

No real employee data is used. By default the workbook is what HR uploads: the sheets
SAP and HC_Teamlease, with the SAP Remarks / Remarks Code cells still empty. HC_Teamlease
has the real headers in the real column positions; only the columns TO CONSIDER needs
are filled.

`with_to_consider=True` gives the file *after* the macro has run: a TO CONSIDER sheet
(with the macro's extra column D) and the SAP Remarks formulas. openpyxl cannot store a
formula together with its cached value, so TO CONSIDER holds the VLOOKUP results as values.

Keyword options let tests create broken or unusual variants.

Command line (from the repo root):  python -m tests.fixtures.make_sample_workbook out.xlsx
"""

import sys
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.core.rules import Thresholds, calculate_status

DEFAULT_START = date(2026, 9, 20)
PUBLIC_HOLIDAY = date(2026, 10, 2)

MANAGERS = {
    "M1": {"id": 20000001, "name": "Fake Manager One", "email": "manager.one@example.com", "status": "Active"},
    "M2": {"id": 20000002, "name": "Fake Manager Two", "email": "Manager.Two@Example.com", "status": "Active"},
}

# (employee id, name, manager key, 5DAG or 6DAG, emp status, lwd, emp email)
EMPLOYEES = [
    (10000001, "Fake Employee A", "M1", "6DAG", "Active", None, "emp.a@example.com"),
    (10000002, "Fake Employee B", "M1", "6DAG", "Active", None, None),
    (10000003, "Fake Employee C", "M1", "5DAG", "Withdrawn", datetime(2026, 10, 10), "emp.c@example.com"),
    (10000004, "Fake Employee D", "M2", "6DAG", "Active", None, "emp.d@example.com"),
    (10000005, "Fake Employee E", "M2", "5DAG", "Active", None, "emp.e@example.com"),
    (10000006, "Fake Employee F", "M2", "6DAG", "Active", None, "emp.f@example.com"),
]
NOT_IN_HC = (10000088, "Fake Only In SAP")
DOJ = datetime(2024, 1, 15)

# Worked-hours pattern, including the rule boundaries.
WORKED_PATTERN = [
    timedelta(hours=9, minutes=8, seconds=8),
    timedelta(hours=4, minutes=59, seconds=59),
    timedelta(hours=5),
    timedelta(hours=7, minutes=59, seconds=59),
    timedelta(hours=8),
    timedelta(0),
    timedelta(hours=9),
]

SAP_HEADERS = [
    "Employee ID", "Employee Name", "Parent Org Unit", "Designation", "Level", "Date",
    "Work Schedule", "Shift Name", "Planned (In)", "Planned (Out)", "Actual (In)", "Actual (Out)",
    "OD (In)", "OD (Out)", "Leave", "Leave Days", "Planned Working Hours", "Worked Hours",
    "Shortfall Hours", "Extra Hours", "Remarks", "Remarks\nCode",
]

# The real HC_Teamlease headers, A to BL.
HC_HEADERS = [
    "Employee ID", "Full Name", "International worker", "Employment Status", "Employment Type", "Level",
    "Joining Date", "Group DOJ", "Date of birth", "Probation Extension Date", "Confirmation Date", "Relieving Date",
    "Reason for Leaving", "Level - DOJ", "Designation - DOJ", "Last Promotion Year", "Serving Notice Period",
    "Consultancy Name", "STA / PDF", "Position ID", "Current Designation", "Role Specification", "Department",
    "Dept. Short Text", "Parent Org Unit", "Operating Unit", "Division", "Function", "Gender",
    "Work schedule (5/6 Days)", "Personnel area", "Personnel Area Text", "Personnel subarea", "Pers. subarea text",
    "Cost Center", "Cost Center Description", "HRBP", "Shift Marker", "Physical Location", "Attendace Location",
    "System ID", "Email ID", "Rep. Mgr Emp No.", "Rep. Manager Name", "RM1 Level", "Rep. Mgr 2 Emp No.",
    "Rep. Manager 2 Name", "RM2 Level", "Section Head Emp No.", "Section Head Name", "Section Head Level",
    "OU Head Employee No.", "OU Head Name", "OU Level", "Marital Status", "Mobile Number", "Highest Qualification",
    "Years of Experience Prior to Syngene", "Years of Experience in Syngene", "Total Years of Experience",
    "Transfer Date", "Reason for Transfer", "Payroll Blocked", "SBP",
]

# The macro's TO CONSIDER, with its column D (a lookup into another workbook).
MACRO_TC_HEADERS = [
    "Employee ID", "Employee Name", "DOJ", None, "Emp Status", "LWD", "Department", "Parent Org",
    "Emp Email ID", "Manager ID", "Manager Name", "Manager Email", "Manager Status",
]
TOTAL_COLUMNS = [
    ("Present", "P"), ("Absent", "A"), ("Weekly Off", "OFF"), ("Comp Off", "CO"),
    ("Maternity Leave", "ML"), ("Maternity Leave", "PL"), ("PH", "PH"), ("Total", None), ("Remarks", None),
]


def default_schedule(base_schedule: str, day: date) -> str:
    """Sunday is OFF for everyone, Saturday is OFF for 5DAG, 2 Oct is a public holiday."""
    if day == PUBLIC_HOLIDAY:
        return "PH"
    if day.weekday() == 6 or (base_schedule == "5DAG" and day.weekday() == 5):
        return "OFF"
    return base_schedule


def remarks_formula(row: int, numeric: bool) -> str:
    """The macro's SAP 'Remarks' (1-4) or 'Remarks Code' (P/A/OFF/PH) formula for one row."""
    p, a, off, ph = ("1", "2", "3", "4") if numeric else ('"P"', '"A"', '"OFF"', '"PH"')
    return (
        f'=IF(G{row}="PH",IF(R{row}>TIME(4,59,59),{p},{ph}),'
        f'IF(OR(G{row}="5DAG",G{row}="6DAG"),IF(R{row}>TIME(4,59,59),{p},{a}),'
        f'IF(G{row}="OFF",IF(R{row}>TIME(7,59,59),{p},{off}))))'
    )


def as_cell_value(worked: timedelta, variant: int) -> object:
    """Store worked hours in the different ways real files do."""
    if variant == 0:
        return (datetime.min + worked).time()                # datetime.time
    if variant == 1:
        return worked                                         # timedelta
    if variant == 2:
        return worked.total_seconds() / 86400                 # Excel day fraction
    total = int(worked.total_seconds())
    return f"{total // 3600:02d}:{total % 3600 // 60:02d}:{total % 60:02d}"  # text


def hc_records(employees: list[list]) -> list[dict[str, object]]:
    """HC_Teamlease rows (by header) for the employees and their managers."""
    records = []
    for emp_id, name, mgr_key, _, status, lwd, email in employees:
        mgr = MANAGERS[mgr_key]
        records.append({
            "Employee ID": emp_id, "Full Name": name, "Employment Status": status, "Joining Date": DOJ,
            "Relieving Date": lwd, "Department": f"Dept {mgr_key}", "Parent Org Unit": "Fake Org", "Email ID": email,
            "Rep. Mgr Emp No.": mgr["id"], "Rep. Manager Name": mgr["name"],
        })
    for key, mgr in MANAGERS.items():
        records.append({
            "Employee ID": mgr["id"], "Full Name": mgr["name"], "Employment Status": mgr["status"],
            "Joining Date": DOJ, "Department": f"Dept {key}", "Parent Org Unit": "Fake Org", "Email ID": mgr["email"],
        })
    return records


def make_sample_workbook(
    path: Path | str,
    *,
    start: date = DEFAULT_START,
    num_days: int = 30,
    include_sap: bool = True,
    include_hc: bool = True,
    with_to_consider: bool = False,
    text_ids: bool = False,
    sap_overrides: dict[tuple[int, date], tuple[str, timedelta]] | None = None,
    drop_sap_rows: set[tuple[int, date]] | None = None,
    employee_not_in_hc: bool = False,
    hc_overrides: dict[int, dict[str, object]] | None = None,
    hc_missing_header: str | None = None,
) -> Path:
    """Write a fake workbook to `path` and return the path.

    hc_overrides: {employee or manager id: {HC_Teamlease header: value}}.
    """
    sap_overrides = sap_overrides or {}
    drop_sap_rows = drop_sap_rows or set()
    hc_overrides = hc_overrides or {}
    dates = [start + timedelta(days=i) for i in range(num_days)]
    thresholds = Thresholds()
    employees = [list(e) for e in EMPLOYEES]
    as_id = (lambda value: str(value)) if text_ids else (lambda value: value)  # noqa: E731 - small local helper

    # 1. Per employee-day: schedule and worked hours.
    sap_days: dict[tuple[int, date], tuple[str, timedelta]] = {}
    for index, (emp_id, _, _, base, *_rest) in enumerate(employees):
        for d_index, day in enumerate(dates):
            sap_days[(emp_id, day)] = (default_schedule(base, day), WORKED_PATTERN[(index + d_index) % len(WORKED_PATTERN)])
    sap_employees = [(e[0], e[1]) for e in employees]
    if employee_not_in_hc:
        sap_employees.append(NOT_IN_HC)
        for day in dates:
            sap_days[(NOT_IN_HC[0], day)] = ("6DAG", timedelta(hours=9))
    sap_days.update(sap_overrides)

    wb = Workbook()
    wb.remove(wb.active)

    # 2. SAP sheet (one row per employee per day).
    if include_sap:
        ws = wb.create_sheet("SAP")
        ws.append(SAP_HEADERS)
        for emp_id, name in sap_employees:
            for d_index, day in enumerate(dates):
                if (emp_id, day) in drop_sap_rows:
                    continue
                schedule, worked = sap_days[(emp_id, day)]
                excel_row = ws.max_row + 1
                ws.append([
                    as_id(emp_id), name, "Fake Org Unit", "Associate", "L1", datetime(day.year, day.month, day.day),
                    schedule, None, time(9), time(18), None, None, None, None, None, None,
                    time(9), as_cell_value(worked, d_index % 4), None, None,
                    # Remarks columns: the macro's formulas (empty in the uploaded file). The app never reads them.
                    remarks_formula(excel_row, numeric=True) if with_to_consider else None,
                    remarks_formula(excel_row, numeric=False) if with_to_consider else None,
                ])

    # 3. HC_Teamlease sheet.
    records = hc_records(employees)
    for record in records:
        record.update(hc_overrides.get(record["Employee ID"], {}))
    if include_hc:
        ws = wb.create_sheet("HC_Teamlease")
        headers = [h for h in HC_HEADERS if h != hc_missing_header]
        ws.append(headers)
        for record in records:
            row = [record.get(header) for header in headers]
            row[0] = as_id(row[0])
            if "Rep. Mgr Emp No." in headers and row[headers.index("Rep. Mgr Emp No.")] is not None:
                row[headers.index("Rep. Mgr Emp No.")] = as_id(row[headers.index("Rep. Mgr Emp No.")])
            ws.append(row)

    # 4. The macro's TO CONSIDER (only for files on which the macro has run).
    if with_to_consider:
        _add_macro_to_consider(wb, sap_employees, records, dates, sap_days, thresholds)
        wb.move_sheet("TO CONSIDER", offset=-(len(wb.sheetnames) - 1))

    path = Path(path)
    wb.save(path)
    return path


def _add_macro_to_consider(wb, sap_employees, records, dates, sap_days, thresholds) -> None:
    ws = wb.create_sheet("TO CONSIDER")
    ws.append(MACRO_TC_HEADERS + [datetime(d.year, d.month, d.day) for d in dates] + [h for h, _ in TOTAL_COLUMNS])
    hc = {record["Employee ID"]: record for record in records}
    first_date_col = len(MACRO_TC_HEADERS) + 1
    span_first, span_last = get_column_letter(first_date_col), get_column_letter(first_date_col + len(dates) - 1)

    for row_number, (emp_id, name) in enumerate(sorted(sap_employees), start=2):
        person = hc.get(emp_id)
        manager = hc.get(person.get("Rep. Mgr Emp No.")) if person else None

        def look(record, header):
            return "#N/A" if record is None else (record.get(header) if record.get(header) is not None else 0)

        row: list[object] = [
            emp_id, name, look(person, "Joining Date"), "#N/A", look(person, "Employment Status"),
            look(person, "Relieving Date"), look(person, "Department"), look(person, "Parent Org Unit"),
            look(person, "Email ID"), look(person, "Rep. Mgr Emp No."), look(person, "Rep. Manager Name"),
            look(manager, "Email ID"), look(manager, "Employment Status"),
        ]
        for day in dates:
            schedule, worked = sap_days[(emp_id, day)]
            row.append(calculate_status(schedule, int(worked.total_seconds()), thresholds))
        span = f"{span_first}{row_number}:{span_last}{row_number}"
        totals = [f'=COUNTIF({span},"{code}")' for _, code in TOTAL_COLUMNS if code]
        first_total = get_column_letter(first_date_col + len(dates))
        last_total = get_column_letter(first_date_col + len(dates) + len(totals) - 1)
        row += totals + [f"=SUM({first_total}{row_number}:{last_total}{row_number})"]
        ws.append(row)


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "sample_workbook.xlsx")
    print(f"Wrote {make_sample_workbook(out)}")
