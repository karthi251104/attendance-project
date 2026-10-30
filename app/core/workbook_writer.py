"""Build the updated workbook for download, doing what the Excel macro does.

Into the original uploaded file (SAP + HC_Teamlease) it:
  1. adds a TO CONSIDER sheet as the first tab, with
       - the headers: A-L, one date per column from M, then the totals and Remarks,
       - one row per employee: A-L as worked out at upload (the macro's VLOOKUPs, stored as values),
       - the day statuses as they are now in the portal (SAP calculation plus managers' edits),
       - the macro's COUNTIF / SUM total formulas,
  2. writes the macro's Remarks / Remarks Code formulas into SAP cells that are empty.
The SAP cells that already hold something and the whole HC_Teamlease sheet are not changed.
"""

import logging
from dataclasses import dataclass
from datetime import datetime

from openpyxl.utils import get_column_letter

from app.core.rules import UNKNOWN_STATUS, Thresholds
from app.core.workbook_reader import (
    DATE_COLUMNS,
    NOT_FOUND,
    REMARKS_HEADER,
    SAP_SHEET,
    TC_COLUMNS,
    TO_CONSIDER_SHEET,
)
from app.core.xlsx_patch import XlsxPatcher, column_index

log = logging.getLogger(__name__)

SHORT_DATE = 14  # Excel's built-in short date format
EXCEL_EPOCH = datetime(1899, 12, 30)
COLUMN_WIDTHS = {"Employee ID": 12, "Employee Name": 28, "DOJ": 11, "Emp Status": 11, "LWD": 11, "Department": 26,
                 "Parent Org": 26, "Emp Email ID": 32, "Manager ID": 12, "Manager Name": 26, "Manager Email": 32,
                 "Manager Status": 14}


class WriterError(Exception):
    pass


@dataclass
class WriteReport:
    employee_rows: int = 0
    day_cells: int = 0
    total_cells: int = 0
    sap_cells: int = 0


def _time_text(seconds: int) -> str:
    """The macro compares with '> TIME(4,59,59)', i.e. one second below the threshold."""
    seconds -= 1
    return f"TIME({seconds // 3600},{seconds % 3600 // 60},{seconds % 60})"


def remarks_formula(row: int, schedule_col: str, worked_col: str, thresholds: Thresholds, numeric: bool) -> str:
    """The macro's SAP formula: Remarks (1=P 2=A 3=OFF 4=PH) or Remarks Code (P/A/OFF/PH)."""
    p, a, off, ph = ("1", "2", "3", "4") if numeric else ('"P"', '"A"', '"OFF"', '"PH"')
    g, r = f"{schedule_col}{row}", f"{worked_col}{row}"
    work, offday, holiday = (_time_text(t) for t in (thresholds.workday_seconds, thresholds.offday_seconds, thresholds.ph_seconds))
    return (f'IF({g}="PH",IF({r}>{holiday},{p},{ph}),'
            f'IF(OR({g}="5DAG",{g}="6DAG"),IF({r}>{work},{p},{a}),'
            f'IF({g}="OFF",IF({r}>{offday},{p},{off}))))')


def excel_serial(value: datetime) -> float:
    """Excel's date number (days since 30/12/1899)."""
    delta = value.replace(tzinfo=None) - EXCEL_EPOCH
    return delta.days + delta.seconds / 86400


def build_updated_workbook(original: bytes, layout: dict, employees: list[dict],
                           thresholds: Thresholds) -> tuple[bytes, WriteReport]:
    report = WriteReport()
    patcher = XlsxPatcher(original)
    if patcher.has_sheet(TO_CONSIDER_SHEET):
        raise WriterError("the uploaded file already has a 'TO CONSIDER' sheet")
    tc = patcher.add_sheet(TO_CONSIDER_SHEET, position=0)  # Sheets.Add puts it before SAP, like the macro
    plain = patcher.add_style(border=True)
    as_date = patcher.add_style(border=True, number_format=SHORT_DATE)
    header = patcher.add_style(border=True, bold=True)
    header_date = patcher.add_style(border=True, bold=True, number_format=SHORT_DATE)

    date_columns: dict[str, str] = layout["date_columns"]
    totals = layout.get("total_columns", [])
    remarks_letter = get_column_letter(column_index(totals[-1]["letter"]) + 1) if totals else None
    employee_rows: dict[str, int] = layout["employee_rows"]
    letters = sorted(date_columns.values(), key=column_index)
    first, last = letters[0], letters[-1]

    # 1. Header row.
    row = layout["header_row"]
    for index, (title, _, _) in enumerate(TC_COLUMNS, start=1):
        _write(tc, f"{get_column_letter(index)}{row}", title, header)
    for iso, letter in date_columns.items():
        _write(tc, f"{letter}{row}", datetime.fromisoformat(iso), header_date)
    for col in totals:
        _write(tc, f"{col['letter']}{row}", col["header"], header)
    if remarks_letter:
        _write(tc, f"{remarks_letter}{row}", REMARKS_HEADER, header)

    # 2. One row per employee, in the order worked out at upload (sorted by Employee ID).
    counted = [col["letter"] for col in totals if col["code"]]
    for emp in sorted(employees, key=lambda e: employee_rows.get(e["employee_id"], 0)):
        row = employee_rows.get(emp["employee_id"])
        if row is None:
            raise WriterError(f"employee {emp['employee_id']} is not in the saved layout")
        report.employee_rows += 1
        for index, ((title, _, _), value) in enumerate(zip(TC_COLUMNS, emp["sheet_cells"]), start=1):
            _write(tc, f"{get_column_letter(index)}{row}", value, as_date if title in DATE_COLUMNS else plain)

        for iso, letter in date_columns.items():
            value = emp["days"].get(iso)
            if value and value != UNKNOWN_STATUS:
                report.day_cells += 1
            else:
                value = None
            _write(tc, f"{letter}{row}", value, plain)

        for col in totals:
            ref = f"{col['letter']}{row}"
            if col["code"]:
                tc.set_formula(ref, f'COUNTIF({first}{row}:{last}{row},"{col["code"]}")')
            elif counted:
                tc.set_formula(ref, f"SUM({counted[0]}{row}:{counted[-1]}{row})")
            tc.set_style(ref, plain)
            report.total_cells += 1
        if remarks_letter:
            tc.set_style(f"{remarks_letter}{row}", plain)

    widths = {get_column_letter(i): COLUMN_WIDTHS[title] for i, (title, _, _) in enumerate(TC_COLUMNS, start=1)}
    widths.update({letter: 10.5 for letter in letters})
    widths.update({col["letter"]: 11 for col in totals})
    tc.set_column_widths(widths)

    # 3. SAP Remarks / Remarks Code formulas (only into empty cells).
    sap_columns: dict[str, str] = layout.get("sap_columns", {})
    if sap_columns.get("work_schedule") and sap_columns.get("worked_hours"):
        sap = patcher.sheet(SAP_SHEET)
        for field_name, numeric in (("remarks", True), ("remarks_code", False)):
            column = sap_columns.get(field_name)
            if not column:
                continue
            for sap_row in range(layout["header_row"] + 1, layout.get("sap_last_row", 0) + 1):
                ref = f"{column}{sap_row}"
                if sap.state(ref).kind == "empty":
                    sap.set_formula(ref, remarks_formula(sap_row, sap_columns["work_schedule"], sap_columns["worked_hours"], thresholds, numeric))
                    report.sap_cells += 1

    patcher.set_full_calc_on_load()
    log.info("Workbook built: employee_rows=%d day_cells=%d total_cells=%d sap_cells=%d",
             report.employee_rows, report.day_cells, report.total_cells, report.sap_cells)
    return patcher.save(), report


def _write(sheet, ref: str, value: object, style: int) -> None:
    """Write one value with the right cell type; None leaves the cell empty (but styled)."""
    if value == NOT_FOUND:
        sheet.set_error(ref, NOT_FOUND)
    elif isinstance(value, datetime):
        sheet.set_number(ref, excel_serial(value))
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        sheet.set_number(ref, value)
    elif value is not None:
        sheet.set_text(ref, str(value))
    sheet.set_style(ref, style)
